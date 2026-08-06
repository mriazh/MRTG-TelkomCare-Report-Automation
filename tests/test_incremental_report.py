import json
import os
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from mrtg_automation.report.excel import ExcelReportGenerator


class TestIncrementalMonthlyReport(unittest.TestCase):
    DATE_01 = "20260801"
    DATE_02 = "20260802"
    TARGET_ID = "TEST-TARGET"

    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.template_path = self.root / "template.xlsx"
        self.output_path = self.root / "monthly-report.xlsx"
        self.mapping_path = self.root / "mapping.txt"
        self.target_list_path = self.root / "targets.csv"
        self.data_dir = self.root / "data"

        template = Workbook()
        template.active.title = "01"
        template.create_sheet("02")
        template["01"]["Z99"] = "existing-template-data"
        template.save(self.template_path)

        self.mapping_path.write_text(
            "\n".join(
                [
                    f"Service Id : {self.TARGET_ID}",
                    "Inbound_Current: A1",
                    "Inbound_Average: A2",
                    "Inbound_Maximum: A3",
                    "Outbound_Current: A4",
                    "Outbound_Average: A5",
                    "Outbound_Maximum: A6",
                    "Image : B1-C1",
                    "",
                ]
            ),
            encoding="utf-8",
        )
        self.target_list_path.write_text(
            f"type,target,ocr_enabled,image_enabled\nSID,{self.TARGET_ID},true,true\n",
            encoding="utf-8",
        )
        self.data_dir.mkdir()
        for date_value in (self.DATE_01, self.DATE_02):
            date_dir = self.data_dir / date_value
            date_dir.mkdir()
            screenshot = date_dir / f"MRTG_{self.TARGET_ID}_{date_value}.png"
            screenshot.write_bytes(b"fake screenshot data" * 128)

        self.config = MagicMock()
        self.config.ocr_max_retries = 1
        self.generator = ExcelReportGenerator(self.config)
        self.ocr_metadata = {
            "values": {
                "Inbound_Current": "6.65 k",
                "Inbound_Average": "99.52 k",
                "Inbound_Maximum": "671.70 k",
                "Outbound_Current": "3.32 k",
                "Outbound_Average": "153.29 k",
                "Outbound_Maximum": "1.82 M",
            },
            "engine_used": "Paddle",
            "decision_reason": "paddle_confident",
            "paddle_confidence": 0.99,
            "paddle_values": {
                "Inbound_Current": "6.65 k",
                "Inbound_Average": "99.52 k",
                "Inbound_Maximum": "671.70 k",
                "Outbound_Current": "3.32 k",
                "Outbound_Average": "153.29 k",
                "Outbound_Maximum": "1.82 M",
            },
            "gemini_values": {},
            "paddle_complete": True,
            "gemini_complete": False,
            "gemini_model": "",
            "gemini_called": False,
            "attempted_models": ["Paddle"],
            "model_failures": {},
            "attempts": 1,
            "gemini_error_reason": None,
        }

    def tearDown(self):
        self.temp_dir.cleanup()

    def _run_date(self, date_value: str):
        with (
            patch("mrtg_automation.report.excel.OCRExtractor._get_engine"),
            patch(
                "mrtg_automation.report.excel.OCRExtractor.extract_mrtg_values_with_metadata",
                return_value=self.ocr_metadata,
            ) as extract,
            patch(
                "mrtg_automation.report.excel.insert_image_to_area",
                return_value=True,
            ) as insert_image,
            patch.dict(os.environ, {"INSERT_IMAGES": "True"}),
        ):
            summary = self.generator.generate(
                "OCR_IMAGE",
                self.data_dir,
                self.template_path,
                self.output_path,
                self.mapping_path,
                self.target_list_path,
                date_filter=date_value,
            )

        return summary, extract, insert_image

    def test_incremental_runs_preserve_prior_sheet_and_audit_entries(self):
        summary_01, extract_01, insert_image_01 = self._run_date(self.DATE_01)

        self.assertTrue(summary_01["success"])
        self.assertEqual(1, summary_01["dates_processed"])
        self.assertEqual(1, summary_01["ocr_ok"])
        extract_01.assert_called_once()
        insert_image_01.assert_called_once()

        workbook_01 = load_workbook(self.output_path, data_only=False)
        self.assertIn("01", workbook_01.sheetnames)
        self.assertEqual("6.65 k", workbook_01["01"]["A1"].value)
        self.assertEqual("existing-template-data", workbook_01["01"]["Z99"].value)
        self.assertIsNone(workbook_01["02"]["A1"].value)

        summary_02, extract_02, insert_image_02 = self._run_date(self.DATE_02)

        self.assertTrue(summary_02["success"])
        self.assertEqual(1, summary_02["dates_processed"])
        self.assertEqual(1, summary_02["ocr_ok"])
        extract_02.assert_called_once()
        insert_image_02.assert_called_once()

        workbook_02 = load_workbook(self.output_path, data_only=False)
        self.assertIn("01", workbook_02.sheetnames)
        self.assertIn("02", workbook_02.sheetnames)
        self.assertEqual("6.65 k", workbook_02["01"]["A1"].value)
        self.assertEqual("existing-template-data", workbook_02["01"]["Z99"].value)
        self.assertEqual("6.65 k", workbook_02["02"]["A1"].value)

        audit_entries = [
            json.loads(line)
            for line in (self.output_path.with_suffix(".ocr-audit.jsonl"))
            .read_text(encoding="utf-8")
            .splitlines()
        ]
        self.assertEqual([self.DATE_01, self.DATE_02], [entry["date"] for entry in audit_entries])
        self.assertEqual(
            self.TARGET_ID,
            audit_entries[0]["target_id"],
        )
        self.assertEqual(
            self.TARGET_ID,
            audit_entries[1]["target_id"],
        )


if __name__ == "__main__":
    unittest.main()
