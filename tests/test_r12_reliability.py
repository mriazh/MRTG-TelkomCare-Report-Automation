import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import MagicMock, patch
import urllib.error

from openpyxl import Workbook

from mrtg_automation.config import Config
from mrtg_automation.report.gemini_ocr import GeminiLegendExtractor
from mrtg_automation.report.ocr import OCRExtractor
from mrtg_automation.report.excel import ExcelReportGenerator


class TestR12ReliabilitySlice(unittest.TestCase):
    def setUp(self):
        self.config = Config()
        self.config.gemini_api_key = "test_key"
        self.config.gemini_models = ["gemini-1", "gemini-2", "gemini-3"]
        self.temp_dir = TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)
        self.dummy_img = self.temp_path / "dummy.png"
        self.dummy_img.write_bytes(b"fake_image_bytes")

    def tearDown(self):
        self.temp_dir.cleanup()

    @patch("urllib.request.urlopen")
    def test_gemini_models_order_and_failed_calls_cap(self, mock_urlopen):
        extractor = GeminiLegendExtractor(self.config)
        mock_urlopen.side_effect = urllib.error.URLError("Connection refused")

        res = extractor.extract_legend(self.dummy_img)

        self.assertFalse(res["complete"])
        self.assertEqual(res["error_reason"], "ALL_MODELS_FAILED")
        self.assertEqual(mock_urlopen.call_count, 9)

    @patch("urllib.request.urlopen")
    def test_gemini_next_model_fallback(self, mock_urlopen):
        def side_effect(req, timeout=30):
            url = req.full_url
            if "gemini-1" in url:
                raise urllib.error.HTTPError(url, 500, "Internal Server Error", {}, None)
            else:
                resp_body = {
                    "candidates": [
                        {
                            "content": {
                                "parts": [
                                    {
                                        "text": "Inbound_Current: 1 k\nInbound_Average: 2 k\nInbound_Maximum: 3 k\nOutbound_Current: 4 k\nOutbound_Average: 5 k\nOutbound_Maximum: 6 k"
                                    }
                                ]
                            }
                        }
                    ]
                }
                resp_data = json.dumps(resp_body).encode("utf-8")
                mock_resp = MagicMock()
                mock_resp.status = 200
                mock_resp.read.return_value = resp_data
                mock_resp.__enter__.return_value = mock_resp
                return mock_resp

        mock_urlopen.side_effect = side_effect

    def test_ocr_finite_retries_and_switching(self):
        ocr = OCRExtractor(self.config)
        ocr._extract_paddle = MagicMock(
            return_value={"values": {}, "confidence": 0.0, "error_reason": "paddle_failed"}
        )

        gemini_success_res = {
            "values": {
                "Inbound_Current": "1 k",
                "Inbound_Average": "2 k",
                "Inbound_Maximum": "3 k",
                "Outbound_Current": "4 k",
                "Outbound_Average": "5 k",
                "Outbound_Maximum": "6 k",
            },
            "model": "gemini-1",
            "complete": True,
            "error_reason": "",
        }
        ocr.gemini_extractor.extract_legend = MagicMock(return_value=gemini_success_res)

        res = ocr.extract(self.dummy_img, max_retries=3)
        self.assertEqual(res["engine_used"], "Gemini")
        self.assertTrue(res["gemini_complete"])

    def test_unresolved_summary_successful_retry_removed_and_safe_fields(self):
        generator = ExcelReportGenerator(self.config)
        data_dir = self.temp_path / "data"
        date_dir = data_dir / "20260901"
        date_dir.mkdir(parents=True)
        img1 = date_dir / "MRTG_target1_20260901.png"
        img2 = date_dir / "MRTG_target2_20260901.png"
        from PIL import Image as PILImg

        PILImg.new("RGB", (10, 10), color="white").save(img1)
        PILImg.new("RGB", (10, 10), color="white").save(img2)

        tpl_path = self.temp_path / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "01"
        ws.cell(row=1, column=1, value="ID")
        ws.cell(row=2, column=1, value="target1")
        ws.cell(row=3, column=1, value="target2")
        wb.save(tpl_path)

        out_path = self.temp_path / "output.xlsx"
        map_path = self.temp_path / "mapping.txt"
        map_path.write_text(
            "Service Id : target1\nInbound_Current: B2\nInbound_Average: B3\nInbound_Maximum: B4\nOutbound_Current: B5\nOutbound_Average: B6\nOutbound_Maximum: B7\nImage : B2-I10\n\n"
            "Service Id : target2\nInbound_Current: C2\nInbound_Average: C3\nInbound_Maximum: C4\nOutbound_Current: C5\nOutbound_Average: C6\nOutbound_Maximum: C7\nImage : C2-I10\n",
            encoding="utf-8",
        )

        list_path = self.temp_path / "list.txt"
        list_path.write_text("1. MRTG : target1\n2. MRTG : target2\n", encoding="utf-8")

        attempt_counter = {"target1": 0, "target2": 0}

        def mock_extract(img_p, max_retries=3):
            t_id = "target1" if "target1" in img_p.name else "target2"
            attempt_counter[t_id] += 1
            if t_id == "target1":
                if attempt_counter[t_id] == 1:
                    return {
                        "values": {},
                        "engine_used": "Paddle",
                        "decision_reason": "paddle_error",
                        "paddle_confidence": 0,
                        "paddle_complete": False,
                        "gemini_complete": False,
                        "gemini_model": "",
                        "gemini_called": True,
                        "paddle_values": {},
                        "gemini_values": {},
                    }
                return {
                    "values": {
                        "Inbound_Current": "1 k",
                        "Inbound_Average": "2 k",
                        "Inbound_Maximum": "3 k",
                        "Outbound_Current": "4 k",
                        "Outbound_Average": "5 k",
                        "Outbound_Maximum": "6 k",
                    },
                    "engine_used": "Paddle",
                    "decision_reason": "paddle_confident",
                    "paddle_confidence": 0.99,
                    "paddle_complete": True,
                    "gemini_complete": False,
                    "gemini_model": "",
                    "gemini_called": False,
                    "paddle_values": {},
                    "gemini_values": {},
                }
            else:
                return {
                    "values": {},
                    "engine_used": "Paddle",
                    "decision_reason": "paddle_error",
                    "paddle_confidence": 0,
                    "paddle_complete": False,
                    "gemini_complete": False,
                    "gemini_model": "",
                    "gemini_called": True,
                    "paddle_values": {},
                    "gemini_values": {},
                }

        with patch("mrtg_automation.report.excel.OCRExtractor") as mock_ocr_cls:
            mock_ocr_cls.extract_mrtg_values_with_metadata.side_effect = mock_extract

            summary = generator.generate(
                report_mode="OCR_IMAGE",
                data_dir=data_dir,
                template_path=tpl_path,
                output_path=out_path,
                mapping_file=map_path,
                list_file=list_path,
            )

        self.assertIn("unresolved_items", summary)
        unresolved = summary["unresolved_items"]
        unresolved_targets = [u["target_id"] for u in unresolved]
        self.assertNotIn("target1", unresolved_targets)
        self.assertIn("target2", unresolved_targets)

        for u in unresolved:
            self.assertIsInstance(u["date"], str)
            self.assertIsInstance(u["target_id"], str)
            self.assertIsInstance(u["mode"], str)
            self.assertIsInstance(u["status"], str)
            self.assertIsInstance(u["error"], str)

    def test_independent_item_continuation(self):
        generator = ExcelReportGenerator(self.config)
        data_dir = self.temp_path / "data"
        date_dir = data_dir / "20260901"
        date_dir.mkdir(parents=True)
        img1 = date_dir / "MRTG_target1_20260901.png"
        img2 = date_dir / "MRTG_target2_20260901.png"
        from PIL import Image as PILImg

        PILImg.new("RGB", (10, 10), color="white").save(img1)
        PILImg.new("RGB", (10, 10), color="white").save(img2)

        tpl_path = self.temp_path / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "01"
        ws.cell(row=1, column=1, value="ID")
        ws.cell(row=2, column=1, value="target1")
        ws.cell(row=3, column=1, value="target2")
        wb.save(tpl_path)

        out_path = self.temp_path / "output.xlsx"
        map_path = self.temp_path / "mapping.txt"
        map_path.write_text(
            "Service Id : target1\nInbound_Current: B2\nInbound_Average: B3\nInbound_Maximum: B4\nOutbound_Current: B5\nOutbound_Average: B6\nOutbound_Maximum: B7\nImage : B2-I10\n\n"
            "Service Id : target2\nInbound_Current: C2\nInbound_Average: C3\nInbound_Maximum: C4\nOutbound_Current: C5\nOutbound_Average: C6\nOutbound_Maximum: C7\nImage : C2-I10\n",
            encoding="utf-8",
        )

        list_path = self.temp_path / "list.txt"
        list_path.write_text("1. MRTG : target1\n2. MRTG : target2\n", encoding="utf-8")

        def mock_extract(img_p, max_retries=3):
            if "target1" in img_p.name:
                raise RuntimeError("Catastrophic error on target1")
            return {
                "values": {
                    "Inbound_Current": "1 k",
                    "Inbound_Average": "2 k",
                    "Inbound_Maximum": "3 k",
                    "Outbound_Current": "4 k",
                    "Outbound_Average": "5 k",
                    "Outbound_Maximum": "6 k",
                },
                "engine_used": "Paddle",
                "decision_reason": "paddle_confident",
                "paddle_confidence": 0.99,
                "paddle_complete": True,
                "gemini_complete": False,
                "gemini_model": "",
                "gemini_called": False,
                "paddle_values": {},
                "gemini_values": {},
            }

        with patch("mrtg_automation.report.excel.OCRExtractor") as mock_ocr_cls:
            mock_ocr_cls.extract_mrtg_values_with_metadata.side_effect = mock_extract

            summary = generator.generate(
                report_mode="OCR_IMAGE",
                data_dir=data_dir,
                template_path=tpl_path,
                output_path=out_path,
                mapping_file=map_path,
                list_file=list_path,
            )

        self.assertTrue(summary["success"])
        self.assertEqual(summary["ocr_ok"], 1)
        self.assertEqual(summary["ocr_fail"], 1)
        self.assertEqual(len(summary["unresolved_items"]), 1)
        self.assertEqual(summary["unresolved_items"][0]["target_id"], "target1")
